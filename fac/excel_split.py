# -*- coding: utf-8 -*-
"""Excel 按人拆分:一个表里有多个人的行时,按行匹配人员,拆成每人一个文件。"""

import os


def split_workbook(path: str, roster, out_tmp_dir: str):
    """按行匹配人员并拆分。

    返回 (results, unmatched_rows):
      results: {Person: 新文件路径},每个新文件 = 各工作表的表头行 + 该人命中的行;
      unmatched_rows: 未匹配到任何人的数据行数(不含表头)。
    解析失败抛异常,由调用方决定按普通文件处理。
    """
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True)

    # sheet -> (header_row, {person: [rows]})
    sheets = {}
    unmatched = 0
    for ws in wb.worksheets:
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue
        header, data = rows[0], rows[1:]
        per = {}
        for row in data:
            text = ' '.join(str(c) for c in row if c is not None)
            persons = roster.match(text)
            if not persons:
                if any(c is not None for c in row):
                    unmatched += 1
                continue
            for p in persons:
                per.setdefault(p, []).append(row)
        if per:
            sheets[ws.title] = (header, per)
    wb.close()

    all_persons = []
    for _title, (_h, per) in sheets.items():
        for p in per:
            if p not in all_persons:
                all_persons.append(p)

    stem, _ext = os.path.splitext(os.path.basename(path))
    results = {}
    for p in all_persons:
        new_wb = openpyxl.Workbook()
        new_wb.remove(new_wb.active)
        for title, (header, per) in sheets.items():
            if p not in per:
                continue
            ws = new_wb.create_sheet(title=title[:31])
            ws.append(list(header))
            for row in per[p]:
                ws.append(list(row))
        out_path = os.path.join(out_tmp_dir, f'{stem}_{p.folder}.xlsx')
        i = 1
        while os.path.exists(out_path):
            out_path = os.path.join(out_tmp_dir, f'{stem}_{p.folder}({i}).xlsx')
            i += 1
        new_wb.save(out_path)
        results[p] = out_path
    return results, unmatched
