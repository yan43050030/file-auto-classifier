# -*- coding: utf-8 -*-
"""反馈核对表:人员 × 来源单位 矩阵,空格即"该单位未反馈",可直接当催办清单。"""

import csv
import os
import datetime

from .util import unique_path


def write_matrix_report(out_dir: str, matrix: dict, log):
    """matrix: {(folder, unit): count}。生成 CSV(Excel 可直接打开);
    装了 openpyxl 时同时生成 xlsx。返回生成的文件路径列表。"""
    if not matrix:
        return []
    units = sorted({u for (_f, u) in matrix})
    folders = []
    for (f, _u) in matrix:
        if f not in folders:
            folders.append(f)
    # 未分类放最后
    folders.sort(key=lambda f: (f == '未分类', f))

    header = ['人员/文件夹'] + units + ['合计']
    rows = []
    for f in folders:
        counts = [matrix.get((f, u), 0) for u in units]
        rows.append([f] + [c if c else '' for c in counts] + [sum(counts)])
    total_row = ['合计'] + \
        [sum(matrix.get((f, u), 0) for f in folders) for u in units] + \
        [sum(matrix.values())]

    written = []
    csv_path = unique_path(out_dir, '反馈核对表.csv')
    with open(csv_path, 'w', newline='', encoding='utf-8-sig') as fp:
        w = csv.writer(fp)
        w.writerow(header)
        w.writerows(rows)
        w.writerow(total_row)
    written.append(csv_path)

    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = '反馈核对表'
        ws.append(header)
        for c in ws[1]:
            c.font = Font(bold=True)
        miss_fill = PatternFill('solid', fgColor='FFF2CC')   # 空格标黄提醒
        for r in rows:
            ws.append(r)
            row_idx = ws.max_row
            for col in range(2, 2 + len(units)):
                if ws.cell(row=row_idx, column=col).value in ('', None):
                    ws.cell(row=row_idx, column=col).fill = miss_fill
        ws.append(total_row)
        for c in ws[ws.max_row]:
            c.font = Font(bold=True)
        ws.column_dimensions['A'].width = 22
        xlsx_path = unique_path(out_dir, '反馈核对表.xlsx')
        wb.save(xlsx_path)
        written.append(xlsx_path)
    except ImportError:
        pass
    except Exception as e:
        log(f'  [提示] 生成 xlsx 核对表失败(已生成 CSV): {e}')
    return written


# ==================== 「文件整理」模式的整理报告 ====================

def write_organize_report(out_dir: str, plan, hs: dict, log):
    """整理报告:总览 + 各文件夹占用 + 类型分布 + 年份分布 + 大文件 Top20。
    返回生成的文件路径列表。"""
    from .health import human_size
    from .filetypes import categorize

    folder_stat, cat_stat, year_stat = {}, {}, {}
    for it in plan:
        rel = it.targets[0] if it.targets else '未分类'
        c, b = folder_stat.get(rel, (0, 0))
        folder_stat[rel] = (c + 1, b + it.size)
        cat = categorize(it.fname)
        c, b = cat_stat.get(cat, (0, 0))
        cat_stat[cat] = (c + 1, b + it.size)
        try:
            y = str(datetime.datetime.fromtimestamp(
                os.path.getmtime(it.src)).year) if os.path.exists(it.src) \
                else '未知'
        except Exception:
            y = '未知'
        c, b = year_stat.get(y, (0, 0))
        year_stat[y] = (c + 1, b + it.size)

    top_large = sorted(plan, key=lambda i: -i.size)[:20]

    sections = [
        ('总览', ['项目', '数值'], [
            ['文件总数', len(plan)],
            ['占用空间', human_size(hs.get('total_bytes', 0))],
            ['重复文件组', hs.get('dup_groups', 0)],
            ['多余的重复文件', hs.get('dup_extra', 0)],
            ['删除重复可省出', human_size(hs.get('dup_bytes', 0))],
            ['垃圾/临时文件', hs.get('junk', 0)],
            ['垃圾文件占用', human_size(hs.get('junk_bytes', 0))],
            ['旧版本文件', hs.get('old_versions', 0)],
        ]),
        ('各文件夹', ['目标文件夹', '文件数', '占用'],
         [[k, v[0], human_size(v[1])]
          for k, v in sorted(folder_stat.items(), key=lambda x: -x[1][1])]),
        ('类型分布', ['类型', '文件数', '占用'],
         [[k, v[0], human_size(v[1])]
          for k, v in sorted(cat_stat.items(), key=lambda x: -x[1][1])]),
        ('年份分布', ['年份', '文件数', '占用'],
         [[k, v[0], human_size(v[1])]
          for k, v in sorted(year_stat.items(), reverse=True)]),
        ('最大的文件', ['文件名', '大小', '归入'],
         [[i.fname, human_size(i.size), i.display_target()]
          for i in top_large]),
    ]

    written = []
    csv_path = unique_path(out_dir, '整理报告.csv')
    with open(csv_path, 'w', newline='', encoding='utf-8-sig') as fp:
        w = csv.writer(fp)
        for title, header, rows in sections:
            w.writerow([f'【{title}】'])
            w.writerow(header)
            w.writerows(rows)
            w.writerow([])
    written.append(csv_path)

    try:
        import openpyxl
        from openpyxl.styles import Font
        wb = openpyxl.Workbook()
        wb.remove(wb.active)
        for title, header, rows in sections:
            ws = wb.create_sheet(title=title[:31])
            ws.append(header)
            for c in ws[1]:
                c.font = Font(bold=True)
            for r in rows:
                ws.append(r)
            ws.column_dimensions['A'].width = 32
            ws.column_dimensions['B'].width = 14
            ws.column_dimensions['C'].width = 16
        xlsx_path = unique_path(out_dir, '整理报告.xlsx')
        wb.save(xlsx_path)
        written.append(xlsx_path)
    except ImportError:
        pass
    except Exception as e:
        log(f'  [提示] 生成 xlsx 整理报告失败(已生成 CSV): {e}')
    return written
