# -*- coding: utf-8 -*-
"""内容级匹配:文件名未命中时,读取文件内容(Excel/Word/PDF/文本)查找人员关键字。

所有解析依赖都是可选的,缺失时对应格式返回空字符串,静默降级。
"""

import re

from .util import read_text_any_encoding

try:
    import openpyxl
    HAS_XLSX = True
except Exception:
    HAS_XLSX = False

try:
    import docx                    # python-docx
    HAS_DOCX = True
except Exception:
    HAS_DOCX = False

try:
    import pdfplumber
    HAS_PDF = True
except Exception:
    HAS_PDF = False

try:
    import xlrd                    # 旧版 .xls(Excel 97-2003)
    HAS_XLS = True
except Exception:
    HAS_XLS = False

try:
    import olefile                 # 旧版 .doc 的 OLE 容器解析(可选,缺失时走原始字节扫描)
    HAS_OLE = True
except Exception:
    HAS_OLE = False

# 内容匹配支持的扩展名(用于界面提示;.doc 为启发式提取,无需额外依赖)
SUPPORTED_EXTS = ['.txt', '.csv', '.doc'] + \
    (['.xlsx', '.xlsm'] if HAS_XLSX else []) + \
    (['.xls'] if HAS_XLS else []) + \
    (['.docx'] if HAS_DOCX else []) + \
    (['.pdf'] if HAS_PDF else [])

_MAX_CHARS = 200_000     # 读取上限,防超大文件拖慢
_MAX_PDF_PAGES = 10


def _read_xlsx(path, limit):
    parts, total = [], 0
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        for ws in wb.worksheets:
            for row in ws.iter_rows(values_only=True):
                for c in row:
                    if c is None:
                        continue
                    s = str(c)
                    parts.append(s)
                    total += len(s)
                    if total >= limit:
                        return ' '.join(parts)
    finally:
        wb.close()
    return ' '.join(parts)


def _read_docx(path, limit):
    d = docx.Document(path)
    parts, total = [], 0
    for p in d.paragraphs:
        if p.text:
            parts.append(p.text)
            total += len(p.text)
            if total >= limit:
                return '\n'.join(parts)
    for t in d.tables:
        for row in t.rows:
            for cell in row.cells:
                if cell.text:
                    parts.append(cell.text)
                    total += len(cell.text)
                    if total >= limit:
                        return '\n'.join(parts)
    return '\n'.join(parts)


def _read_pdf(path, limit):
    parts, total = [], 0
    with pdfplumber.open(path) as pdf:
        for page in pdf.pages[:_MAX_PDF_PAGES]:
            t = page.extract_text() or ''
            if t:
                parts.append(t)
                total += len(t)
                if total >= limit:
                    break
    return '\n'.join(parts)


def _cell_str(c):
    """单元格值转字符串;整数浮点去掉 .0,避免证号/账号变形。"""
    if isinstance(c, float) and c.is_integer():
        return str(int(c))
    return str(c)


def _read_xls(path, limit):
    """旧版 .xls:xlrd 逐表逐行读取。"""
    parts, total = [], 0
    wb = xlrd.open_workbook(path)
    for ws in wb.sheets():
        for r in range(ws.nrows):
            for c in ws.row_values(r):
                if c is None or c == '':
                    continue
                s = _cell_str(c)
                parts.append(s)
                total += len(s)
                if total >= limit:
                    return ' '.join(parts)
    return ' '.join(parts)


_DOC_TEXT_RE = re.compile(r'[一-鿿0-9Xx]{2,}')
_DOC_MAX_BYTES = 4 << 20     # 最多读 4MB,防超大文件拖慢


def _read_doc(path, limit):
    """旧版 .doc 的启发式文本提取。

    不做 Word 二进制格式的完整解析——匹配姓名/证号只需要"内容里有没有
    这个字串",所以:优先取 OLE 容器里的 WordDocument 流(装了 olefile 时),
    否则退回读原始字节;然后分别按 UTF-16LE 和 GBK 宽松解码,扫出所有
    连续的中文/数字片段拼成文本。会混入少量乱码,但不影响子串匹配。"""
    data = None
    if HAS_OLE:
        try:
            ole = olefile.OleFileIO(path)
            try:
                if ole.exists('WordDocument'):
                    data = ole.openstream('WordDocument').read()
            finally:
                ole.close()
        except Exception:
            data = None
    if data is None:
        with open(path, 'rb') as f:
            data = f.read(_DOC_MAX_BYTES)
    data = data[:_DOC_MAX_BYTES]

    parts, total = [], 0
    # UTF-16LE 按两种字节偏移各扫一遍(文本片段起点不一定偶数对齐)
    for decoded in (data.decode('utf-16-le', errors='ignore'),
                    data[1:].decode('utf-16-le', errors='ignore'),
                    data.decode('gbk', errors='ignore')):
        for m in _DOC_TEXT_RE.finditer(decoded):
            parts.append(m.group(0))
            total += len(m.group(0))
            if total >= limit:
                return ' '.join(parts)
    return ' '.join(parts)


def extract_text(path: str, max_chars: int = _MAX_CHARS) -> str:
    """提取文件文本内容用于匹配。不支持/解析失败返回空字符串。"""
    low = path.lower()
    try:
        if low.endswith(('.txt', '.csv')):
            return (read_text_any_encoding(path) or '')[:max_chars]
        if low.endswith(('.xlsx', '.xlsm')) and HAS_XLSX:
            return _read_xlsx(path, max_chars)
        if low.endswith('.docx') and HAS_DOCX:
            return _read_docx(path, max_chars)
        if low.endswith('.pdf') and HAS_PDF:
            return _read_pdf(path, max_chars)
        if low.endswith('.xls') and HAS_XLS:
            return _read_xls(path, max_chars)
        if low.endswith('.doc'):
            return _read_doc(path, max_chars)
    except Exception:
        return ''
    return ''


def can_match(path: str) -> bool:
    return any(path.lower().endswith(e) for e in SUPPORTED_EXTS)
